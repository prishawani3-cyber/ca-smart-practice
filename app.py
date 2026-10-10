import sqlite3
from pathlib import Path
from datetime import date, datetime
import io

import pandas as pd
import plotly.express as px
import streamlit as st

# =========================================================
# CA SMART PRACTICE - SINGLE FILE PRACTICE APPLICATION
# =========================================================
st.set_page_config(page_title="CA Smart Practice", page_icon="📊", layout="wide")
BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "ca_practice.db"

# ------------------------- DATABASE -------------------------
def connect():
    return sqlite3.connect(DB_PATH)

def init_db():
    with connect() as conn:
        conn.execute("""
        CREATE TABLE IF NOT EXISTS clients(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL, pan TEXT DEFAULT '', gstin TEXT DEFAULT '',
            mobile TEXT DEFAULT '', email TEXT DEFAULT '',
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS audit_tasks(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client TEXT DEFAULT '', task TEXT NOT NULL,
            owner TEXT DEFAULT '', due_date TEXT DEFAULT '',
            status TEXT DEFAULT 'Pending', notes TEXT DEFAULT ''
        )""")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS compliance_tasks(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client TEXT DEFAULT '', compliance TEXT NOT NULL,
            due_date TEXT NOT NULL, status TEXT DEFAULT 'Pending',
            notes TEXT DEFAULT ''
        )""")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS tax_records(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client TEXT DEFAULT '', financial_year TEXT DEFAULT '',
            record_type TEXT NOT NULL, description TEXT DEFAULT '',
            amount REAL DEFAULT 0, notes TEXT DEFAULT ''
        )""")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS tds_records(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client TEXT DEFAULT '', payee TEXT DEFAULT '',
            section TEXT DEFAULT '', payment_date TEXT DEFAULT '',
            payment_amount REAL DEFAULT 0, rate REAL DEFAULT 0,
            threshold REAL DEFAULT 0, notes TEXT DEFAULT ''
        )""")

def query_df(sql, params=()):
    with connect() as conn:
        return pd.read_sql_query(sql, conn, params=params)

def execute(sql, params=()):
    with connect() as conn:
        conn.execute(sql, params)

def csv_download_button(df, label, filename):
    st.download_button(label, df.to_csv(index=False).encode("utf-8-sig"),
                       file_name=filename, mime="text/csv")

init_db()

# ------------------------- HELPERS -------------------------
def read_upload(file):
    if file.name.lower().endswith(".csv"):
        return pd.read_csv(file)
    return pd.read_excel(file)

def clean_columns(df):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    return df

def find_column(df, candidates):
    lookup = {str(c).strip().lower(): c for c in df.columns}
    for candidate in candidates:
        if candidate.lower() in lookup:
            return lookup[candidate.lower()]
    return None

def money(value):
    return f"₹{value:,.2f}"

def clients_list():
    return query_df("SELECT name FROM clients ORDER BY name")["name"].tolist()

# ------------------------- NAVIGATION -------------------------
st.title("📊 CA Smart Practice")
st.caption("All-in-One CA Practice Assistant | Practice version")
pages = [
    "Dashboard", "Client Management", "Accounting & Excel Tools",
    "Bank Reconciliation", "GST Reconciliation", "TDS Tools",
    "Income Tax Organizer", "Audit Management", "Compliance Calendar",
    "Reports & Exports"
]
st.sidebar.title("CA Practice")
page = st.sidebar.radio("Select Module", pages)
st.sidebar.divider()
st.sidebar.caption("Practice version — verify all tax outputs before use.")

# =========================================================
# DASHBOARD
# =========================================================
if page == "Dashboard":
    st.header("Practice Dashboard")
    client_count = query_df("SELECT COUNT(*) n FROM clients")["n"].iloc[0]
    pending_audit = query_df("SELECT COUNT(*) n FROM audit_tasks WHERE status != 'Completed'")["n"].iloc[0]
    pending_compliance = query_df("SELECT COUNT(*) n FROM compliance_tasks WHERE status != 'Completed'")["n"].iloc[0]
    c1, c2, c3 = st.columns(3)
    c1.metric("Total Clients", int(client_count))
    c2.metric("Open Audit Tasks", int(pending_audit))
    c3.metric("Open Compliance Tasks", int(pending_compliance))
    task_data = pd.DataFrame({"Category": ["Audit tasks", "Compliance tasks"],
                              "Open": [int(pending_audit), int(pending_compliance)]})
    st.plotly_chart(px.bar(task_data, x="Category", y="Open", title="Open Work Overview"),
                    use_container_width=True)
    st.info("Start by adding clients, then use the modules from the left sidebar.")

# =========================================================
# CLIENT MANAGEMENT
# =========================================================
elif page == "Client Management":
    st.header("👥 Client Management")
    tab1, tab2, tab3 = st.tabs(["Add Client", "Search / Export", "Edit / Delete"])
    with tab1:
        with st.form("add_client", clear_on_submit=True):
            name = st.text_input("Client / Business Name *")
            pan = st.text_input("PAN").strip().upper()
            gstin = st.text_input("GSTIN").strip().upper()
            mobile = st.text_input("Mobile")
            email = st.text_input("Email")
            save = st.form_submit_button("Save Client")
            if save:
                if not name.strip():
                    st.error("Client name is required.")
                elif pan and len(pan) != 10:
                    st.error("PAN should contain 10 characters.")
                elif gstin and len(gstin) != 15:
                    st.error("GSTIN should contain 15 characters.")
                else:
                    execute("INSERT INTO clients(name,pan,gstin,mobile,email) VALUES(?,?,?,?,?)",
                            (name.strip(), pan, gstin, mobile.strip(), email.strip()))
                    st.success("Client saved.")
                    st.rerun()
    with tab2:
        search = st.text_input("Search client")
        df = query_df("""SELECT id AS 'Client ID', name AS 'Client Name',
                         pan AS PAN, gstin AS GSTIN, mobile AS Mobile, email AS Email
                         FROM clients WHERE name LIKE ? OR pan LIKE ? OR gstin LIKE ?
                         OR mobile LIKE ? OR email LIKE ? ORDER BY id DESC""",
                      tuple([f"%{search}%"] * 5))
        st.dataframe(df, use_container_width=True, hide_index=True)
        csv_download_button(df, "Download Client Register", "client_register.csv")
    with tab3:
        clients = query_df("SELECT * FROM clients ORDER BY name")
        if clients.empty:
            st.info("No clients registered.")
        else:
            labels = {f'{r["id"]} - {r["name"]}': int(r["id"]) for _, r in clients.iterrows()}
            chosen = st.selectbox("Choose client", list(labels.keys()))
            cid = labels[chosen]
            row = query_df("SELECT * FROM clients WHERE id=?", (cid,)).iloc[0]
            with st.form("edit_client"):
                name2 = st.text_input("Name", str(row["name"]))
                pan2 = st.text_input("PAN", str(row["pan"] or ""))
                gstin2 = st.text_input("GSTIN", str(row["gstin"] or ""))
                mobile2 = st.text_input("Mobile", str(row["mobile"] or ""))
                email2 = st.text_input("Email", str(row["email"] or ""))
                update = st.form_submit_button("Update client")
                if update:
                    if not name2.strip():
                        st.error("Name is required.")
                    elif pan2 and len(pan2.strip()) != 10:
                        st.error("PAN should contain 10 characters.")
                    elif gstin2 and len(gstin2.strip()) != 15:
                        st.error("GSTIN should contain 15 characters.")
                    else:
                        execute("UPDATE clients SET name=?,pan=?,gstin=?,mobile=?,email=? WHERE id=?",
                                (name2.strip(), pan2.strip().upper(), gstin2.strip().upper(),
                                 mobile2.strip(), email2.strip(), cid))
                        st.success("Client updated.")
                        st.rerun()
            if st.checkbox("Confirm permanent deletion", key=f"confirm_del_{cid}"):
                if st.button("Delete selected client", type="secondary"):
                    execute("DELETE FROM clients WHERE id=?", (cid,))
                    st.success("Client deleted.")
                    st.rerun()

# =========================================================
# ACCOUNTING & EXCEL TOOLS
# =========================================================
elif page == "Accounting & Excel Tools":
    st.header("📁 Accounting & Excel Tools")
    file = st.file_uploader("Upload Excel or CSV", type=["xlsx", "xls", "csv"], key="accounting_file")
    if file:
        try:
            df = read_upload(file)
            st.success(f"Loaded {len(df):,} rows and {len(df.columns)} columns.")
            st.dataframe(df, use_container_width=True)
            info = pd.DataFrame({"Column": df.columns,
                                 "Data type": [str(df[c].dtype) for c in df.columns],
                                 "Missing values": [int(df[c].isna().sum()) for c in df.columns]})
            st.subheader("Column quality")
            st.dataframe(info, use_container_width=True, hide_index=True)
            csv_download_button(df, "Download as CSV", "accounting_export.csv")
        except Exception as e:
            st.error(f"Unable to read file: {e}")

# =========================================================
# BANK RECONCILIATION
# =========================================================
elif page == "Bank Reconciliation":
    st.header("🏦 Bank Reconciliation")
    st.caption("Matches unique references and compares transaction amounts.")
    bank_file = st.file_uploader("Upload Bank Statement", type=["xlsx", "xls", "csv"], key="bank_file")
    book_file = st.file_uploader("Upload Cash Book", type=["xlsx", "xls", "csv"], key="cash_file")
    if bank_file and book_file:
        try:
            bank, book = clean_columns(read_upload(bank_file)), clean_columns(read_upload(book_file))
            bref = find_column(bank, ["Reference", "Ref", "Transaction ID", "UTR", "Cheque No"])
            cref = find_column(book, ["Reference", "Ref", "Transaction ID", "UTR", "Cheque No"])
            bamt = find_column(bank, ["Amount", "Transaction Amount", "Debit/Credit"])
            camt = find_column(book, ["Amount", "Transaction Amount", "Debit/Credit"])
            if not all([bref, cref, bamt, camt]):
                st.error("Both files need a reference column (Reference/Ref/UTR) and an Amount column.")
            else:
                bank = bank.rename(columns={bref: "Reference", bamt: "Bank Amount"})
                book = book.rename(columns={cref: "Reference", camt: "Cash Book Amount"})
                for d in (bank, book):
                    d["Reference"] = d["Reference"].fillna("").astype(str).str.strip()
                bank["Bank Amount"] = pd.to_numeric(bank["Bank Amount"], errors="coerce")
                book["Cash Book Amount"] = pd.to_numeric(book["Cash Book Amount"], errors="coerce")
                if bank["Reference"].eq("").any() or book["Reference"].eq("").any():
                    st.error("Blank references found. Please correct them.")
                elif bank["Reference"].duplicated().any() or book["Reference"].duplicated().any():
                    st.error("Duplicate references found. Resolve them before using this simple matching method.")
                elif bank["Bank Amount"].isna().any() or book["Cash Book Amount"].isna().any():
                    st.error("Invalid or blank amounts found.")
                else:
                    result = bank.merge(book, on="Reference", how="outer", indicator=True, suffixes=("_Bank", "_Book"), validate="one_to_one")
                    result["Status"] = result.apply(
                        lambda r: "Only in Bank" if r["_merge"] == "left_only"
                        else ("Only in Cash Book" if r["_merge"] == "right_only"
                              else ("Matched" if abs(r["Bank Amount"] - r["Cash Book Amount"]) < 0.01
                                    else "Amount Mismatch")), axis=1)
                    counts = result["Status"].value_counts()
                    cols = st.columns(4)
                    for col, label in zip(cols, ["Matched", "Only in Bank", "Only in Cash Book", "Amount Mismatch"]):
                        col.metric(label, int(counts.get(label, 0)))
                    output = result.drop(columns="_merge")
                    st.dataframe(output.style.apply(
                        lambda row: ["background-color: " + (
                            "#d4edda" if row["Status"] == "Matched" else
                            "#fff3cd" if row["Status"] == "Only in Bank" else
                            "#f8d7da" if row["Status"] == "Only in Cash Book" else "#f5c6cb"
                        )] * len(row), axis=1), use_container_width=True, hide_index=True)
                    csv_download_button(output, "Download reconciliation report", "bank_reconciliation.csv")
        except Exception as e:
            st.error(f"Bank reconciliation error: {e}")
    else:
        st.info("Upload both files to reconcile.")

# =========================================================
# GST RECONCILIATION
# =========================================================
elif page == "GST Reconciliation":
    st.header("🧾 GST Reconciliation — Purchase Register vs GSTR-2B")
    st.caption("Upload both files to compare supplier GSTIN, invoice number, taxable value and tax.")

    left, right = st.columns(2)
    with left:
        pr_file = st.file_uploader(
            "📁 Purchase Register Excel/CSV",
            type=["xlsx", "xls", "csv"], key="gst_pr"
        )
    with right:
        g2b_file = st.file_uploader(
            "📁 GSTR-2B Excel/CSV",
            type=["xlsx", "xls", "csv"], key="gst_2b"
        )

    tolerance = st.number_input(
        "Difference tolerance (₹)", min_value=0.0,
        value=1.0, step=0.5, key="gst_tolerance"
    )

    def gst_map(df):
        return {
            "GSTIN": find_column(df, [
                "GSTIN", "Supplier GSTIN", "GSTIN of Supplier",
                "GSTIN/UIN of Supplier", "GSTIN of supplier"
            ]),
            "Invoice": find_column(df, [
                "Invoice No", "Invoice Number", "Bill No",
                "Document Number", "Invoice No.", "Document No",
                "Invoice number"
            ]),
            "Date": find_column(df, [
                "Invoice Date", "Date", "Document Date"
            ]),
            "Taxable": find_column(df, [
                "Taxable Value", "Taxable Amount", "Taxable"
            ]),
            "IGST": find_column(df, ["IGST", "Integrated Tax"]),
            "CGST": find_column(df, ["CGST", "Central Tax"]),
            "SGST": find_column(df, [
                "SGST", "UTGST", "State Tax"
            ]),
            "TotalTax": find_column(df, [
                "Total Tax", "Tax Amount", "Total GST", "Total Tax Amount"
            ])
        }

    def gst_number(series):
        return pd.to_numeric(
            series.astype(str)
            .str.replace(",", "", regex=False)
            .str.replace("₹", "", regex=False)
            .str.strip(),
            errors="coerce"
        ).fillna(0.0)

    def gst_normalize(df, mapping, suffix):
        out = pd.DataFrame()
        out["GSTIN"] = (
            df[mapping["GSTIN"]].fillna("").astype(str).str.upper().str.strip()
            if mapping["GSTIN"] else ""
        )
        out["Invoice"] = (
            df[mapping["Invoice"]].fillna("").astype(str)
            .str.upper().str.strip().str.replace(r"\\s+", "", regex=True)
            if mapping["Invoice"] else ""
        )
        out["Invoice"] = out["Invoice"].str.replace(r"\\.0$", "", regex=True)
        out["GSTIN"] = out["GSTIN"].replace({"NAN": "", "NONE": ""})
        out["Date"] = (
            df[mapping["Date"]].astype(str)
            if mapping["Date"] else ""
        )
        for field in ["Taxable", "IGST", "CGST", "SGST"]:
            out[field] = (
                gst_number(df[mapping[field]])
                if mapping[field] else 0.0
            )
        out["TotalTax"] = (
            gst_number(df[mapping["TotalTax"]])
            if mapping["TotalTax"]
            else out["IGST"] + out["CGST"] + out["SGST"]
        )
        rename = {
            field: f"{field}_{suffix}"
            for field in ["Date", "Taxable", "IGST", "CGST", "SGST", "TotalTax"]
        }
        return out.rename(columns=rename)

    if not pr_file or not g2b_file:
        st.info("Upload both files above to start reconciliation.")
    else:
        try:
            pr = clean_columns(read_upload(pr_file))
            g2b = clean_columns(read_upload(g2b_file))
            pm, gm = gst_map(pr), gst_map(g2b)

            missing = []
            for label, mapping in [
                ("Purchase Register", pm), ("GSTR-2B", gm)
            ]:
                if not mapping["GSTIN"]:
                    missing.append(f"{label}: supplier GSTIN column")
                if not mapping["Invoice"]:
                    missing.append(f"{label}: invoice number column")

            if missing:
                st.error(
                    "Could not identify these required columns: "
                    + "; ".join(missing)
                    + ". Rename the relevant Excel headings and upload again."
                )
                st.write("Purchase Register columns:", list(pr.columns))
                st.write("GSTR-2B columns:", list(g2b.columns))
            else:
                a = gst_normalize(pr, pm, "PR")
                b = gst_normalize(g2b, gm, "2B")
                keys = ["GSTIN", "Invoice"]

                if a[keys].eq("").any(axis=1).any() or b[keys].eq("").any(axis=1).any():
                    st.error(
                        "Blank GSTIN or invoice number found. Correct blank keys "
                        "in the source files before reconciling."
                    )
                elif a.duplicated(keys).any() or b.duplicated(keys).any():
                    st.error(
                        "Duplicate GSTIN + invoice combinations found. "
                        "Please review duplicates in the source files first. "
                        "No reconciliation was performed, to avoid incorrect matches."
                    )
                    d1, d2 = st.columns(2)
                    with d1:
                        st.write("Purchase Register duplicate keys")
                        st.dataframe(
                            a[a.duplicated(keys, keep=False)],
                            use_container_width=True, hide_index=True
                        )
                    with d2:
                        st.write("GSTR-2B duplicate keys")
                        st.dataframe(
                            b[b.duplicated(keys, keep=False)],
                            use_container_width=True, hide_index=True
                        )
                else:
                    result = a.merge(
                        b, on=keys, how="outer", indicator=True,
                        validate="one_to_one"
                    )

                    compare_fields = [
                        "Taxable", "IGST", "CGST", "SGST", "TotalTax"
                    ]

                    def gst_status(row):
                        if row["_merge"] == "left_only":
                            return "Missing in GSTR-2B"
                        if row["_merge"] == "right_only":
                            return "Missing in Purchase Register"
                        for field in compare_fields:
                            x = float(row.get(f"{field}_PR", 0) or 0)
                            y = float(row.get(f"{field}_2B", 0) or 0)
                            if abs(x - y) > tolerance:
                                return "Value/Tax Mismatch"
                        return "Matched"

                    result["Status"] = result.apply(gst_status, axis=1)
                    output = result.drop(columns=["_merge"])
                    counts = output["Status"].value_counts()

                    st.subheader("📊 Reconciliation Summary")
                    c1, c2, c3, c4 = st.columns(4)
                    c1.metric("✅ Matched", int(counts.get("Matched", 0)))
                    c2.metric(
                        "❌ Missing in 2B",
                        int(counts.get("Missing in GSTR-2B", 0))
                    )
                    c3.metric(
                        "⚠️ Missing in Purchase Register",
                        int(counts.get("Missing in Purchase Register", 0))
                    )
                    c4.metric(
                        "🔎 Value/Tax Mismatch",
                        int(counts.get("Value/Tax Mismatch", 0))
                    )

                    def gst_colour(row):
                        status = row.get("Status", "")
                        colours = {
                            "Matched": "#D4EDDA",
                            "Missing in GSTR-2B": "#F8D7DA",
                            "Missing in Purchase Register": "#FFF3CD",
                            "Value/Tax Mismatch": "#F5C6CB"
                        }
                        return [
                            f"background-color: {colours.get(status, '#FFFFFF')}"
                        ] * len(row)

                    st.subheader("📋 Detailed Results")
                    st.dataframe(
                        output.style.apply(gst_colour, axis=1).format({
                            col: "{:,.2f}" for col in output.columns
                            if col.startswith((
                                "Taxable_", "IGST_", "CGST_",
                                "SGST_", "TotalTax_"
                            ))
                        }),
                        use_container_width=True, hide_index=True
                    )

                    # Export a formatted Excel workbook
                    excel_buffer = io.BytesIO()
                    try:
                        from openpyxl.styles import Font, PatternFill, Alignment
                        from openpyxl.utils import get_column_letter

                        with pd.ExcelWriter(
                            excel_buffer, engine="openpyxl"
                        ) as writer:
                            output.to_excel(
                                writer, index=False,
                                sheet_name="GST Reconciliation"
                            )
                            ws = writer.sheets["GST Reconciliation"]
                            header_fill = PatternFill(
                                fill_type="solid", fgColor="17365D"
                            )
                            for cell in ws[1]:
                                cell.fill = header_fill
                                cell.font = Font(
                                    bold=True, color="FFFFFF"
                                )
                                cell.alignment = Alignment(
                                    horizontal="center"
                                )

                            status_col = list(output.columns).index("Status") + 1
                            status_fills = {
                                "Matched": "D4EDDA",
                                "Missing in GSTR-2B": "F8D7DA",
                                "Missing in Purchase Register": "FFF3CD",
                                "Value/Tax Mismatch": "F5C6CB"
                            }
                            for row_num in range(2, ws.max_row + 1):
                                status = ws.cell(row_num, status_col).value
                                fill = PatternFill(
                                    fill_type="solid",
                                    fgColor=status_fills.get(status, "FFFFFF")
                                )
                                for cell in ws[row_num]:
                                    cell.fill = fill

                            ws.freeze_panes = "A2"
                            ws.auto_filter.ref = ws.dimensions
                            for cells in ws.columns:
                                letter = get_column_letter(cells[0].column)
                                longest = max(
                                    len(str(cell.value or "")) for cell in cells
                                )
                                ws.column_dimensions[letter].width = min(
                                    max(longest + 2, 12), 36
                                )

                        st.download_button(
                            "⬇️ Download Formatted Excel Report",
                            data=excel_buffer.getvalue(),
                            file_name="GST_Reconciliation_Report.xlsx",
                            mime=(
                                "application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.sheet"
                            )
                        )
                    except ImportError:
                        st.error(
                            "Excel export requires openpyxl. Install it using: "
                            "py -m pip install openpyxl"
                        )

                    st.caption(
                        "Practice tool only. Review invoice dates, credit/debit "
                        "notes, amendments, ITC eligibility and applicable GST "
                        "rules manually before finalising."
                    )
        except Exception as e:
            st.error(f"GST reconciliation error: {e}")

# =========================================================
# TDS TOOLS
# =========================================================
elif page == "TDS Tools":
    st.header("🧮 TDS Tools — Upload Your Excel File")
    st.write(
        "Upload your own Excel/CSV file or download the ready-made FY 2026-27 template. "
        "The app adds the applicable Income-tax Act reference based on the credit/payment date "
        "and nature of payment, then calculates using the rate and threshold in your file."
    )

    # Ready-to-fill Excel template for tax year 2026-27.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    template_buffer = io.BytesIO()
    template_wb = Workbook()
    template_ws = template_wb.active
    template_ws.title = "TDS Working FY 2026-27"
    template_headers = [
        "Client Name", "Payee / Vendor", "PAN", "Payee Residency",
        "Nature of Payment", "Credit/Payment Date", "Payment Amount (₹)",
        "TDS Rate (%)", "Threshold (₹)", "TDS Base (₹)",
        "TDS Amount (₹)", "Net Payment (₹)", "Notes"
    ]
    template_ws.append(template_headers)
    template_ws.append([
        "Example Client", "Example Vendor", "ABCDE1234F", "Resident",
        "Contractor / work contract", "2026-04-05", 10000, 1, 30000,
        None, None, None, "Example only — replace with actual data"
    ])
    for r in range(2, 502):
        template_ws.cell(r, 10, f'=IF(OR(G{r}="",H{r}=""),"",IF(G{r}>I{r},G{r},0))')
        template_ws.cell(r, 11, f'=IF(OR(J{r}="",H{r}=""),"",ROUND(J{r}*H{r}/100,2))')
        template_ws.cell(r, 12, f'=IF(OR(G{r}="",K{r}=""),"",ROUND(G{r}-K{r},2))')
    nature_list = (
        'Salary,Contractor / work contract,Professional services,Technical services,'
        'Commission / brokerage,Rent - land/building/furniture,Rent - machinery/equipment,'
        'Interest on securities,Interest other than securities,Partner payment,'
        'Purchase of goods,Immovable property transfer,Non-resident payment,Other / verify'
    )
    dv = DataValidation(type="list", formula1='"' + nature_list + '"', allow_blank=True)
    template_ws.add_data_validation(dv)
    dv.add("E2:E501")
    dv_res = DataValidation(type="list", formula1='"Resident,Non-resident"', allow_blank=True)
    template_ws.add_data_validation(dv_res)
    dv_res.add("D2:D501")
    header_fill = PatternFill("solid", fgColor="17365D")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    formula_fill = PatternFill("solid", fgColor="DDEBF7")
    thin_side = Side(style="thin", color="D9E1F2")
    for cell in template_ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(2, 502):
        for c in range(1, 10):
            template_ws.cell(r, c).fill = input_fill
        for c in range(10, 13):
            template_ws.cell(r, c).fill = formula_fill
        for c in range(1, 14):
            template_ws.cell(r, c).border = Border(bottom=thin_side)
        for c in [7, 9, 10, 11, 12]:
            template_ws.cell(r, c).number_format = '₹#,##0.00'
        template_ws.cell(r, 8).number_format = '0.00'
    widths = [22, 24, 16, 18, 34, 21, 20, 15, 18, 18, 18, 18, 38]
    for i, width in enumerate(widths, 1):
        template_ws.column_dimensions[get_column_letter(i)].width = width
    template_ws.freeze_panes = "A2"
    template_ws.auto_filter.ref = "A1:M501"
    template_ws.row_dimensions[1].height = 36
    template_wb.calculation.fullCalcOnLoad = True
    template_wb.calculation.forceFullCalc = True
    template_wb.calculation.calcMode = "auto"
    template_wb.save(template_buffer)
    st.download_button(
        "⬇️ Download Ready TDS Excel Template — Income-tax Act 2025",
        data=template_buffer.getvalue(),
        file_name="TDS_Working_Income_Tax_Act_2025_FY_2026_27.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    st.caption("Yellow columns are inputs; blue columns contain formulas. Rates/thresholds are editable and must be checked for each transaction.")
    st.divider()

    uploaded_tds = st.file_uploader(
        "📤 Upload your TDS Excel/CSV file",
        type=["xlsx", "xls", "csv"],
        key="tds_bulk_upload"
    )

    st.info(
        "Your file must contain a payment amount column and a TDS rate column. "
        "A threshold column is optional. Column headings can be like "
        "'Payment Amount', 'Amount', 'TDS Rate (%)', 'Rate', and 'Threshold'."
    )

    if uploaded_tds is not None:
        try:
            if uploaded_tds.name.lower().endswith(".csv"):
                bulk = pd.read_csv(uploaded_tds)
            else:
                bulk = pd.read_excel(uploaded_tds)

            if bulk.empty:
                st.error("The uploaded file has no data rows.")
            else:
                bulk.columns = [str(c).strip() for c in bulk.columns]
                lower_map = {str(c).strip().lower(): c for c in bulk.columns}

                def locate(names):
                    for name in names:
                        if name.lower() in lower_map:
                            return lower_map[name.lower()]
                    return None

                amount_col = locate([
                    "Payment Amount (₹)", "Payment Amount", "Amount",
                    "Payment amount", "Gross Amount", "Credit Amount",
                    "Invoice Amount", "Bill Amount"
                ])
                rate_col = locate([
                    "TDS Rate (%)", "TDS Rate", "Rate (%)", "Rate",
                    "TDS %", "TDS Percentage"
                ])
                threshold_col = locate([
                    "Threshold (₹)", "Threshold", "Threshold Amount",
                    "Threshold Amount (₹)"
                ])

                if amount_col is None or rate_col is None:
                    st.error(
                        "Required columns not found. Your Excel file needs "
                        "columns for payment amount and TDS rate. "
                        "For example: 'Payment Amount' and 'TDS Rate (%)'."
                    )
                    st.write("Columns detected in your file:")
                    st.write(list(bulk.columns))
                    st.markdown(
                        "Example: `Payee | Payment Amount | TDS Rate (%) | Threshold`"
                    )
                else:
                    calc_df = bulk.copy()

                    # Income-tax Act, 2025 section reference helper.
                    # The Act governing TDS is determined by the earlier of credit/payment.
                    date_col = locate([
                        "Credit/Payment Date", "Credit Date", "Payment Date",
                        "Transaction Date", "Date"
                    ])
                    nature_col = locate([
                        "Nature of Payment", "Payment Type", "Expense Type",
                        "TDS Nature", "Transaction Type"
                    ])
                    residency_col = locate(["Payee Residency", "Residency", "Resident / Non-resident", "Payee Status"])

                    nature_options = [
                        "Salary", "Contractor / work contract", "Professional services",
                        "Technical services", "Commission / brokerage",
                        "Rent - land/building/furniture", "Rent - machinery/equipment",
                        "Interest on securities", "Interest other than securities",
                        "Partner payment", "Purchase of goods", "Immovable property transfer",
                        "Non-resident payment", "Other / verify"
                    ]
                    if nature_col is None:
                        st.info(
                            "Optional: choose the nature of payment to add a section reference. "
                            "For row-by-row section tagging, include a column named 'Nature of Payment'."
                        )
                        chosen_nature = st.selectbox(
                            "Nature of payment (used when your file has no nature column)",
                            nature_options, key="tds_nature_choice"
                        )
                    else:
                        chosen_nature = None

                    def tds_section_for(nature, transaction_date, residency="Resident"):
                        nature_text = str(nature or "").strip().lower()
                        resident = str(residency or "resident").strip().lower()
                        old_map = [
                            ("salary", "192"),
                            ("contract", "194C"), ("work contract", "194C"),
                            ("professional", "194J"), ("technical", "194J"),
                            ("commission", "194H"), ("brokerage", "194H"),
                            ("rent", "194I"), ("interest on securities", "193"),
                            ("interest", "194A"), ("partner", "194T"),
                            ("purchase of goods", "194Q"),
                            ("immovable property", "194-IA"),
                            ("non-resident", "195"), ("non resident", "195")
                        ]
                        # Only well-supported common references are assigned automatically.
                        new_map = [
                            ("salary", "Section 392 — Salaries"),
                            ("contract", "Section 393(1), Table for residents, Sl. No. 6(i) / 6(ii) — contractor/work; verify payer/payee conditions"),
                            ("work contract", "Section 393(1), Table for residents, Sl. No. 6(i) / 6(ii) — contractor/work; verify payer/payee conditions"),
                            ("professional", "Section 393(1), Table for residents, Sl. No. 6(iii) — professional services; verify specified-person condition"),
                            ("technical", "Section 393(1), Table for residents, Sl. No. 6(iii) — technical services"),
                            ("commission", "Section 393(1), Table for residents, Sl. No. 1(ii) or 6(ii) — verify payer category"),
                            ("brokerage", "Section 393(1), Table for residents, Sl. No. 1(ii) or 6(ii) — verify payer category"),
                            ("rent - machinery", "Section 393(1), Table for residents, Sl. No. 2(i) or 2(ii) — machinery/equipment; verify payer category"),
                            ("rent - land", "Section 393(1), Table for residents, Sl. No. 2(i) or 2(ii) — land/building; verify payer category"),
                            ("rent", "Section 393(1), Table for residents, Sl. No. 2(i) / 2(ii) — verify payer category"),
                            ("interest on securities", "Section 393(1), Table for residents, Sl. No. 5(i)"),
                            ("interest other than securities", "Section 393(1), Table for residents, Sl. No. 5(ii) / 5(iii) — verify payer/payee conditions"),
                            ("interest", "Section 393(1), Table for residents, Sl. No. 5(i) / 5(ii) / 5(iii) — verify exact category"),
                            ("partner", "Section 393(3), Table for any person, Sl. No. 7 — partner salary/remuneration/commission/interest"),
                            ("purchase of goods", "Section 393(1), Table for residents, Sl. No. 8(ii) — purchase of goods; verify threshold/other conditions"),
                            ("immovable property", "Section 393(1), Table for residents, Sl. No. 3(i) — property transfer; verify consideration/stamp value and conditions"),
                            ("non-resident", "Section 393(2), applicable non-resident Table item — determine income type/treaty/rate"),
                            ("non resident", "Section 393(2), applicable non-resident Table item — determine income type/treaty/rate")
                        ]
                        if pd.isna(transaction_date) or str(transaction_date).strip() == "":
                            return "Date missing — Act not determined", "Enter credit/payment date and verify section"
                        parsed_date = pd.to_datetime(transaction_date, errors="coerce", dayfirst=True)
                        if pd.isna(parsed_date):
                            return "Invalid date — Act not determined", "Check date and verify section"
                        if parsed_date.date() < date(2026, 4, 1):
                            section = next((sec for key, sec in old_map if key in nature_text), "to be verified")
                            return "Income-tax Act, 1961", "Section " + section
                        if "non-resident" in nature_text or "non resident" in nature_text or resident == "non-resident":
                            return "Income-tax Act, 2025", "Section 393(2) — select the applicable non-resident Table item"
                        section = next((desc for key, desc in new_map if key in nature_text),
                                       "Section 393 — select the correct Table and serial number manually")
                        return "Income-tax Act, 2025", section

                    nature_values = (calc_df[nature_col] if nature_col is not None
                                     else pd.Series([chosen_nature] * len(calc_df), index=calc_df.index))
                    if date_col is not None:
                        date_values = calc_df[date_col]
                    else:
                        st.warning(
                            "No credit/payment date column detected. Add 'Credit Date' or "
                            "'Payment Date' to let the app select the applicable Act."
                        )
                        date_values = pd.Series([None] * len(calc_df), index=calc_df.index)
                    residency_values = (calc_df[residency_col] if residency_col is not None
                                        else pd.Series(["Resident"] * len(calc_df), index=calc_df.index))
                    section_results = [
                        tds_section_for(nature, txn_date, residency)
                        for nature, txn_date, residency in zip(nature_values, date_values, residency_values)
                    ]
                    calc_df["Applicable Act"] = [x[0] for x in section_results]
                    calc_df["TDS Section Reference"] = [x[1] for x in section_results]

                    def to_number(series):
                        return pd.to_numeric(
                            series.astype(str)
                            .str.replace(",", "", regex=False)
                            .str.replace("₹", "", regex=False)
                            .str.replace("%", "", regex=False)
                            .str.strip()
                            .replace({"": None, "nan": None, "None": None}),
                            errors="coerce"
                        )

                    calc_df["Payment Amount (₹)"] = to_number(
                        calc_df[amount_col]
                    )
                    calc_df["TDS Rate (%)"] = to_number(
                        calc_df[rate_col]
                    )

                    if threshold_col is not None:
                        calc_df["Threshold (₹)"] = to_number(
                            calc_df[threshold_col]
                        ).fillna(0)
                    else:
                        calc_df["Threshold (₹)"] = 0.0

                    invalid = (
                        calc_df["Payment Amount (₹)"].isna()
                        | calc_df["TDS Rate (%)"].isna()
                    )
                    if invalid.any():
                        st.warning(
                            f"{int(invalid.sum())} row(s) have a missing or "
                            "invalid payment amount/rate. Those rows are marked "
                            "for review and their calculated amounts are left blank."
                        )

                    calc_df["TDS Base (₹)"] = (
                        calc_df["Payment Amount (₹)"].where(
                            calc_df["Payment Amount (₹)"]
                            > calc_df["Threshold (₹)"], 0
                        )
                    )
                    calc_df["TDS Amount (₹)"] = (
                        calc_df["TDS Base (₹)"]
                        * calc_df["TDS Rate (%)"] / 100
                    ).round(2)
                    calc_df["Net Payment (₹)"] = (
                        calc_df["Payment Amount (₹)"]
                        - calc_df["TDS Amount (₹)"]
                    ).round(2)

                    if invalid.any():
                        calc_df.loc[
                            invalid,
                            [
                                "TDS Base (₹)", "TDS Amount (₹)",
                                "Net Payment (₹)"
                            ]
                        ] = None
                    calc_df["Calculation Status"] = invalid.map(
                        {True: "Check amount/rate", False: "Calculated"}
                    )

                    total_payment = calc_df["Payment Amount (₹)"].sum()
                    total_tds = calc_df["TDS Amount (₹)"].sum()
                    total_net = calc_df["Net Payment (₹)"].sum()

                    st.subheader("📊 Calculation Summary")
                    m1, m2, m3, m4 = st.columns(4)
                    m1.metric("Rows uploaded", f"{len(calc_df):,}")
                    m2.metric("Total payment (₹)", f"{total_payment:,.2f}")
                    m3.metric("Total TDS (₹)", f"{total_tds:,.2f}")
                    m4.metric("Net payment (₹)", f"{total_net:,.2f}")

                    st.subheader("Calculated TDS — all rows")
                    st.dataframe(
                        calc_df.style.format({
                            "Payment Amount (₹)": "{:,.2f}",
                            "TDS Rate (%)": "{:,.2f}",
                            "Threshold (₹)": "{:,.2f}",
                            "TDS Base (₹)": "{:,.2f}",
                            "TDS Amount (₹)": "{:,.2f}",
                            "Net Payment (₹)": "{:,.2f}"
                        }, na_rep="—"),
                        use_container_width=True,
                        hide_index=True
                    )

                    from openpyxl.styles import Font, PatternFill, Alignment
                    from openpyxl.utils import get_column_letter

                    out_buffer = io.BytesIO()
                    with pd.ExcelWriter(
                        out_buffer, engine="openpyxl"
                    ) as writer:
                        calc_df.to_excel(
                            writer, index=False, sheet_name="TDS Calculations"
                        )
                        result_ws = writer.sheets["TDS Calculations"]
                        header_fill = PatternFill(
                            "solid", fgColor="17365D"
                        )
                        input_fill = PatternFill(
                            "solid", fgColor="FFF2CC"
                        )
                        calc_fill = PatternFill(
                            "solid", fgColor="DDEBF7"
                        )
                        for cell in result_ws[1]:
                            cell.fill = header_fill
                            cell.font = Font(
                                bold=True, color="FFFFFF"
                            )
                            cell.alignment = Alignment(
                                horizontal="center", wrap_text=True
                            )
                        for row_num in range(2, result_ws.max_row + 1):
                            for col_num in range(1, result_ws.max_column + 1):
                                cell = result_ws.cell(row_num, col_num)
                                if cell.value is not None:
                                    cell.fill = calc_fill
                        result_ws.freeze_panes = "A2"
                        result_ws.auto_filter.ref = result_ws.dimensions
                        for cells in result_ws.columns:
                            max_len = min(
                                max(len(str(c.value or "")) for c in cells) + 2,
                                38
                            )
                            result_ws.column_dimensions[
                                get_column_letter(cells[0].column)
                            ].width = max_len
                    st.download_button(
                        "⬇️ Download Calculated TDS Excel",
                        data=out_buffer.getvalue(),
                        file_name="TDS_Calculated_Report.xlsx",
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet"
                        ),
                        key="tds_bulk_download"
                    )

        except Exception as e:
            st.error(f"Could not process TDS file: {e}")

    st.warning(
        "Important: for payments/credits up to 31 March 2026, the Income-tax Act, 1961 generally applies; "
        "from 1 April 2026, salary TDS is under section 392 and other TDS is generally under section 393 "
        "of the Income-tax Act, 2025. The app provides a section-reference aid for common categories, "
        "not a final legal determination. Rates and thresholds are taken from your uploaded file; "
        "verify exact Table/serial number, payer status, payee residency, PAN, aggregation, exemptions, "
        "special provisions and current Rules before filing or payment."
    )

    st.divider()
    st.subheader("Saved TDS working")
    tds_df = query_df("SELECT * FROM tds_records ORDER BY id DESC")
    if tds_df.empty:
        st.caption("No saved TDS workings yet.")
    else:
        st.dataframe(tds_df, use_container_width=True, hide_index=True)
        csv_download_button(tds_df, "Export TDS working", "tds_working.csv")

# =========================================================
# INCOME TAX ORGANIZER
# =========================================================
elif page == "Income Tax Organizer":
    st.header("📑 Income Tax Organizer")
    st.write("Maintain income and deduction items for a client. This is an organizer, not a return-filing system.")
    with st.form("tax_record_form", clear_on_submit=True):
        client = st.selectbox("Client (optional)", [""] + clients_list())
        fy = st.selectbox("Financial Year", ["2024-25", "2025-26", "2026-27"])
        kind = st.selectbox("Record type", ["Salary", "House Property", "Business/Profession", "Capital Gains", "Other Income", "Chapter VI-A Deduction", "TDS / Advance Tax", "Other"])
        description = st.text_input("Description / document reference")
        amount = st.number_input("Amount (₹)", min_value=0.0, value=0.0, step=1000.0)
        notes = st.text_input("Notes")
        save = st.form_submit_button("Save tax record")
        if save:
            execute("INSERT INTO tax_records(client,financial_year,record_type,description,amount,notes) VALUES(?,?,?,?,?,?)",
                    (client, fy, kind, description, amount, notes))
            st.success("Tax record saved.")
            st.rerun()
    tax_df = query_df("SELECT * FROM tax_records ORDER BY id DESC")
    if not tax_df.empty:
        f1, f2 = st.columns(2)
        chosen_fy = f1.selectbox("Filter financial year", ["All"] + sorted(tax_df["financial_year"].dropna().unique().tolist()))
        chosen_client = f2.selectbox("Filter client", ["All"] + sorted([x for x in tax_df["client"].dropna().unique().tolist() if x]))
        view = tax_df.copy()
        if chosen_fy != "All": view = view[view["financial_year"] == chosen_fy]
        if chosen_client != "All": view = view[view["client"] == chosen_client]
        st.dataframe(view, use_container_width=True, hide_index=True)
        income_types = ["Salary", "House Property", "Business/Profession", "Capital Gains", "Other Income"]
        income_total = view[view["record_type"].isin(income_types)]["amount"].sum()
        deductions = view[view["record_type"] == "Chapter VI-A Deduction"]["amount"].sum()
        taxes_paid = view[view["record_type"] == "TDS / Advance Tax"]["amount"].sum()
        a, b, c = st.columns(3)
        a.metric("Recorded income items", money(income_total))
        b.metric("Recorded deductions", money(deductions))
        c.metric("TDS / advance tax entered", money(taxes_paid))
        csv_download_button(view, "Export tax organizer", "income_tax_organizer.csv")
    else:
        st.info("No tax records yet. Add records above.")

# =========================================================
# AUDIT MANAGEMENT
# =========================================================
elif page == "Audit Management":
    st.header("🗂️ Audit Management")
    st.write("Create an audit checklist and track work status.")
    with st.form("audit_task_form", clear_on_submit=True):
        client = st.selectbox("Client (optional)", [""] + clients_list())
        task = st.text_input("Audit task / checklist item *")
        owner = st.text_input("Assigned to")
        due = st.date_input("Due date", value=date.today())
        status = st.selectbox("Status", ["Pending", "In Progress", "Completed", "Blocked"])
        notes = st.text_area("Working paper reference / notes")
        save = st.form_submit_button("Add audit task")
        if save:
            if not task.strip():
                st.error("Task description is required.")
            else:
                execute("INSERT INTO audit_tasks(client,task,owner,due_date,status,notes) VALUES(?,?,?,?,?,?)",
                        (client, task.strip(), owner.strip(), str(due), status, notes.strip()))
                st.success("Audit task saved.")
                st.rerun()
    audit = query_df("SELECT * FROM audit_tasks ORDER BY due_date, id")
    if not audit.empty:
        st.dataframe(audit, use_container_width=True, hide_index=True)
        csv_download_button(audit, "Export audit checklist", "audit_checklist.csv")
        task_ids = {f'{r["id"]} - {r["task"]}': int(r["id"]) for _, r in audit.iterrows()}
        with st.expander("Update audit task status"):
            label = st.selectbox("Select task", list(task_ids.keys()))
            new_status = st.selectbox("New status", ["Pending", "In Progress", "Completed", "Blocked"])
            if st.button("Update status"):
                execute("UPDATE audit_tasks SET status=? WHERE id=?", (new_status, task_ids[label]))
                st.success("Status updated.")
                st.rerun()
        with st.expander("Delete audit task"):
            label_del = st.selectbox("Select task to delete", list(task_ids.keys()), key="audit_delete")
            if st.button("Delete task"):
                execute("DELETE FROM audit_tasks WHERE id=?", (task_ids[label_del],))
                st.success("Task deleted.")
                st.rerun()
    else:
        st.info("No audit tasks have been created.")

# =========================================================
# COMPLIANCE CALENDAR
# =========================================================
elif page == "Compliance Calendar":
    st.header("📅 Compliance Calendar")
    st.write("Add due dates, track status and export your task list. Confirm statutory dates from official sources.")
    with st.form("compliance_form", clear_on_submit=True):
        client = st.selectbox("Client (optional)", [""] + clients_list())
        compliance = st.text_input("Compliance / filing name *", placeholder="Example: GST return review")
        due = st.date_input("Due date", value=date.today(), key="compliance_due")
        status = st.selectbox("Status", ["Pending", "In Progress", "Completed"])
        notes = st.text_input("Notes")
        save = st.form_submit_button("Add compliance")
        if save:
            if not compliance.strip():
                st.error("Compliance name is required.")
            else:
                execute("INSERT INTO compliance_tasks(client,compliance,due_date,status,notes) VALUES(?,?,?,?,?)",
                        (client, compliance.strip(), str(due), status, notes.strip()))
                st.success("Compliance item saved.")
                st.rerun()
    tasks = query_df("SELECT * FROM compliance_tasks ORDER BY due_date")
    if not tasks.empty:
        tasks["Due date"] = pd.to_datetime(tasks["due_date"], errors="coerce").dt.date
        today = date.today()
        overdue_count = int(((tasks["Due date"] < today) & (tasks["status"] != "Completed")).sum())
        pending_count = int((tasks["status"] != "Completed").sum())
        c1, c2 = st.columns(2)
        c1.metric("Open items", pending_count)
        c2.metric("Overdue items", overdue_count)
        st.dataframe(tasks, use_container_width=True, hide_index=True)
        csv_download_button(tasks, "Export compliance calendar", "compliance_calendar.csv")
        task_ids = {f'{r["id"]} - {r["compliance"]}': int(r["id"]) for _, r in tasks.iterrows()}
        with st.expander("Update compliance status"):
            label = st.selectbox("Select compliance item", list(task_ids.keys()))
            new_status = st.selectbox("Set status", ["Pending", "In Progress", "Completed"])
            if st.button("Save compliance status"):
                execute("UPDATE compliance_tasks SET status=? WHERE id=?", (new_status, task_ids[label]))
                st.success("Compliance status updated.")
                st.rerun()
    else:
        st.info("No compliance deadlines entered yet.")

# =========================================================
# REPORTS & EXPORTS
# =========================================================
elif page == "Reports & Exports":
    st.header("📤 Reports & Exports")
    st.write("Download registers and working papers created in this app.")
    report_options = {
        "Client Register": ("SELECT id, name, pan, gstin, mobile, email, created_at FROM clients ORDER BY name", "client_register.csv"),
        "Audit Checklist": ("SELECT * FROM audit_tasks ORDER BY due_date", "audit_checklist.csv"),
        "Compliance Calendar": ("SELECT * FROM compliance_tasks ORDER BY due_date", "compliance_calendar.csv"),
        "TDS Working": ("SELECT * FROM tds_records ORDER BY id DESC", "tds_working.csv"),
        "Income Tax Organizer": ("SELECT * FROM tax_records ORDER BY id DESC", "income_tax_organizer.csv")
    }
    selected = st.selectbox("Choose report", list(report_options.keys()))
    sql, filename = report_options[selected]
    report = query_df(sql)
    st.metric("Rows in report", len(report))
    st.dataframe(report, use_container_width=True, hide_index=True)
    csv_download_button(report, f"Download {selected}", filename)
    summary = pd.DataFrame([
        {"Module": "Clients", "Records": len(query_df("SELECT id FROM clients"))},
        {"Module": "Audit tasks", "Records": len(query_df("SELECT id FROM audit_tasks"))},
        {"Module": "Compliance items", "Records": len(query_df("SELECT id FROM compliance_tasks"))},
        {"Module": "TDS working", "Records": len(query_df("SELECT id FROM tds_records"))},
        {"Module": "Income tax records", "Records": len(query_df("SELECT id FROM tax_records"))}
    ])
    st.subheader("Application record summary")
    st.dataframe(summary, use_container_width=True, hide_index=True)
    csv_download_button(summary, "Download module summary", "ca_smart_practice_summary.csv")

st.sidebar.caption("Data is stored locally in ca_practice.db")
